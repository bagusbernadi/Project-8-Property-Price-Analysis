"""Cek silang Q14 dengan spreadsheet: workbook berisi rumus MEDIAN(IF(...)), dihitung ulang oleh LibreOffice Calc
(pengganti Excel, karena Excel tidak tersedia). Hasil dibandingkan dengan median dari DuckDB."""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from common import *  # noqa: E402
import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

con = duckdb.connect(str(DB_PATH), read_only=True)
v = con.execute("SELECT city, district, price_per_m2 FROM listings WHERE flag_ppm2_out = 0 ORDER BY ads_id").df()
sql_city = con.execute("SELECT city, median(price_per_m2) m FROM listings WHERE flag_ppm2_out=0 GROUP BY city").df().set_index("city").m
sql_dist = con.execute(f"SELECT city, district, median(price_per_m2) m, count(*) n FROM listings WHERE flag_ppm2_out=0 "
                       f"GROUP BY city, district HAVING count(*) >= {MIN_DISTRICT_N}").df()
con.close()
n = len(v); last = n + 1

wb = Workbook()
data = wb.active; data.title = "data"
data.append(["city", "district", "price_per_m2"])
for r in v.itertuples(index=False):
    data.append([r.city, r.district, float(r.price_per_m2)])
kota = wb.create_sheet("kota"); kota.append(["city", "median_excel", "n_excel"])
for i, c in enumerate(sorted(sql_city.index), start=2):
    kota[f"A{i}"] = c
    kota[f"B{i}"] = ArrayFormula(f"B{i}", f"=MEDIAN(IF(data!$A$2:$A${last}=A{i},data!$C$2:$C${last}))")
    kota[f"C{i}"] = f"=COUNTIF(data!$A$2:$A${last},A{i})"
kec = wb.create_sheet("kecamatan"); kec.append(["city", "district", "median_excel", "n_excel"])
for i, r in enumerate(sql_dist.sort_values(["city", "district"]).itertuples(index=False), start=2):
    kec[f"A{i}"], kec[f"B{i}"] = r.city, r.district
    kec[f"C{i}"] = ArrayFormula(f"C{i}", f"=MEDIAN(IF((data!$A$2:$A${last}=A{i})*(data!$B$2:$B${last}=B{i}),data!$C$2:$C${last}))")
    kec[f"D{i}"] = f"=COUNTIFS(data!$A$2:$A${last},A{i},data!$B$2:$B${last},B{i})"
src = RES_ANALYSIS / "q14_excel_crosscheck.xlsx"
wb.save(src)

tmp = Path(tempfile.mkdtemp())
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp), str(src)],
               check=True, capture_output=True, timeout=600)
calc = load_workbook(tmp / src.name, data_only=True)
rows = []
kc = calc["kota"]
xk = {kc[f"A{r}"].value: (kc[f"B{r}"].value, kc[f"C{r}"].value) for r in range(2, kc.max_row + 1)}
d_ok = sum(1 for c, m in sql_city.items() if abs(xk[c][0] - m) < 1e-6)
rows.append(("kota", len(sql_city), d_ok, max(abs(xk[c][0] - m) for c, m in sql_city.items())))
kk = calc["kecamatan"]
xd = {(kk[f"A{r}"].value, kk[f"B{r}"].value): (kk[f"C{r}"].value, kk[f"D{r}"].value) for r in range(2, kk.max_row + 1)}
ok_d = [abs(xd[(r.city, r.district)][0] - r.m) < 1e-6 and xd[(r.city, r.district)][1] == r.n for r in sql_dist.itertuples()]
rows.append(("kecamatan", len(sql_dist), sum(ok_d), max(abs(xd[(r.city, r.district)][0] - r.m) for r in sql_dist.itertuples())))
out = pd.DataFrame(rows, columns=["level", "jumlah_dibandingkan", "sama", "selisih_maks"])
out["status"] = ["PASS" if a == b else "FAIL" for a, b in zip(out.jumlah_dibandingkan, out.sama)]
out.to_csv(RES_ANALYSIS / "05_crosscheck_spreadsheet.csv", index=False)
shutil.rmtree(tmp, ignore_errors=True)
print(out.to_string(index=False))
